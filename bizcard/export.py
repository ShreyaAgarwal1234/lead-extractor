"""Excel (.xlsx) export."""
from __future__ import annotations

import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MAX_COL_WIDTH = 45


def write_excel(frame: pd.DataFrame, directory: str | Path | None = None) -> str:
    """Write the lead table to a styled .xlsx file and return its path."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(list(frame.columns))

    for row in frame.itertuples(index=False):
        ws.append(list(row))
    # Security: text read from a card must never be executed as an Excel formula.
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"

    for cell in ws[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        longest = max(len(str(c.value)) if c.value is not None else 0 for c in column)
        ws.column_dimensions[column[0].column_letter].width = min(longest + 3, _MAX_COL_WIDTH)

    out_dir = Path(directory) if directory else Path(tempfile.mkdtemp(prefix="leads_"))
    path = out_dir / f"leads_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb.save(path)
    return str(path)
