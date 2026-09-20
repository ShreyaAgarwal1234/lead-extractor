import pandas as pd
from openpyxl import load_workbook

from bizcard.export import write_excel
from bizcard.schema import LEAD_COLUMNS


def test_excel_roundtrip(tmp_path):
    frame = pd.DataFrame([["Rahul", "Sharma", "Manager", "ABC", "Delhi", "+91 98765 43210", "r@abc.com", "1.png"]],
                         columns=LEAD_COLUMNS)
    path = write_excel(frame, tmp_path)

    ws = load_workbook(path).active
    assert [c.value for c in ws[1]] == LEAD_COLUMNS
    assert [c.value for c in ws[2]][:2] == ["Rahul", "Sharma"]
    assert ws["F2"].value == "+91 98765 43210"  # phone stays text, keeps '+'
    assert ws.freeze_panes == "A2"


def test_formula_text_is_not_executed(tmp_path):
    frame = pd.DataFrame([["=1+1", "", "", "", "", "", "", "x.png"]], columns=LEAD_COLUMNS)
    ws = load_workbook(write_excel(frame, tmp_path)).active
    assert ws["A2"].data_type == "s"  # stored as string, not as formula


def test_empty_frame_still_writes_header(tmp_path):
    ws = load_workbook(write_excel(pd.DataFrame(columns=LEAD_COLUMNS), tmp_path)).active
    assert [c.value for c in ws[1]] == LEAD_COLUMNS
